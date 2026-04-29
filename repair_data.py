import sqlite3
import pandas as pd
import os
import openpyxl
from sqlalchemy import text

def scrub_database(db_path):
    """
    Removes duplicates from the local SQLite staging.db tables.
    Uses 'rowid' to keep only the first occurrence of unique merchant/month pairs.
    """
    if not os.path.exists(db_path):
        return False

    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    # Tables to scrub and their unique keys
    scrub_targets = {
        "PROCESSED_CARD_MONTHLY": ["MERCHANT_GROUP", "MERCHANT_BRAND", "TRANSACTION_MONTH"],
        "PROCESSED_CARD_HISTORY": ["MERCHANT_GROUP", "MERCHANT_BRAND", "TRANSACTION_MONTH"],
        "PROCESSED_CARD_SHARE":   ["MERCHANT_GROUP", "MERCHANT_BRAND", "TRANSACTION_MONTH"],
        "TARGET":                 ["MERCHANT_GROUP", "PM"]
    }

    try:
        for table, keys in scrub_targets.items():
            # Check if table exists
            cursor.execute(f"SELECT name FROM sqlite_master WHERE type='table' AND name='{table}'")
            if not cursor.fetchone():
                continue

            # Build deduplication query
            keys_str = ", ".join(keys)
            delete_sql = f"""
                DELETE FROM {table}
                WHERE rowid NOT IN (
                    SELECT MIN(rowid)
                    FROM {table}
                    GROUP BY {keys_str}
                )
            """
            cursor.execute(delete_sql)
        
        # ── Specific Normalization: Yoshinoya March 2025 Spike (7.1M -> ~141k) ──
        # This fix is critical for dashboard scale
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='PROCESSED_CARD_MONTHLY'")
        if cursor.fetchone():
            cursor.execute("""
                UPDATE PROCESSED_CARD_MONTHLY
                SET TRX_QRIS_OFFUS = TRX_QRIS_OFFUS / 50
                WHERE MERCHANT_GROUP = 'YOSHINOYA' 
                  AND TRANSACTION_MONTH = '2025-03'
                  AND TRX_QRIS_OFFUS > 5000000
            """)

        conn.commit()
    finally:
        conn.close()
    return True

def scrub_staging_tables(db_path):
    """
    Removes duplicates from the three raw staging tables in staging.db.
    Uses natural business keys (not EDW_FETCH_DATE) so re-fetched data
    with a new fetch timestamp is collapsed to the first occurrence.
    """
    if not os.path.exists(db_path):
        return False

    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    staging_targets = {
        "ALL_MID":        ["MERCHANT_ID", "TERMINAL_ID"],
        "CARD_SHARE":     ["MERCHANT_GROUP", "MERCHANT_BRAND", "TRANSACTION_MONTH"],
        "WEEKLY_MONITOR": ["MERCHANT_GROUP", "YEAR", "WEEK_NUM"],
    }

    try:
        for table, keys in staging_targets.items():
            cursor.execute(
                f"SELECT name FROM sqlite_master WHERE type='table' AND name='{table}'"
            )
            if not cursor.fetchone():
                continue

            keys_str = ", ".join(keys)
            cursor.execute(f"""
                DELETE FROM {table}
                WHERE rowid NOT IN (
                    SELECT MIN(rowid)
                    FROM {table}
                    GROUP BY {keys_str}
                )
            """)

        conn.commit()
    finally:
        conn.close()
    return True


def scrub_staging_neon(engine, schema="public"):
    """
    Removes duplicates from the three raw staging tables in Neon PostgreSQL.
    Uses ctid to keep only the first occurrence per natural business key.
    Returns a dict of {table: message}.
    """
    staging_targets = {
        "all_mid":        ["merchant_id", "terminal_id"],
        "card_share":     ["merchant_group", "merchant_brand", "transaction_month"],
        "weekly_monitor": ["merchant_group", "year", "week_num"],
    }

    from sqlalchemy import text as _text
    results = {}
    with engine.begin() as conn:
        for table, keys in staging_targets.items():
            full_table = f'"{schema}"."{table}"'
            keys_str = ", ".join(keys)
            try:
                res = conn.execute(_text(f"""
                    DELETE FROM {full_table}
                    WHERE ctid NOT IN (
                        SELECT MIN(ctid)
                        FROM {full_table}
                        GROUP BY {keys_str}
                    )
                """))
                results[table] = f"Removed {res.rowcount} duplicate(s)."
            except Exception as e:
                results[table] = f"Error: {e}"
    return results


def scrub_excel_card_share(path):
    """Removes duplicates from master_card_share.xlsx 'Realisasi' sheet."""
    if not os.path.exists(path):
        return False

    # Load and de-duplicate
    df = pd.read_excel(path, sheet_name=None)
    if "Realisasi" in df:
        cols = df["Realisasi"].columns.tolist()
        # Essential keys for uniqueness
        keys = ["MERCHANT_GROUP", "MERCHANT_BRAND", "TRANSACTION_MONTH"]
        keys = [k for k in keys if k in cols]
        
        df["Realisasi"] = df["Realisasi"].drop_duplicates(subset=keys, keep='first')
        
        # Apply Yoshinoya fix in Excel too
        if 'MERCHANT_GROUP' in cols and 'TRANSACTION_MONTH' in cols and 'TRX_QRIS_OFFUS' in cols:
            mask = (df["Realisasi"]['MERCHANT_GROUP'] == 'YOSHINOYA') & \
                   (df["Realisasi"]['TRANSACTION_MONTH'].astype(str).str.contains('2025-03')) & \
                   (df["Realisasi"]['TRX_QRIS_OFFUS'] > 5000000)
            df["Realisasi"].loc[mask, 'TRX_QRIS_OFFUS'] = df["Realisasi"].loc[mask, 'TRX_QRIS_OFFUS'] / 50

        # Save back
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            for sheet_name, sheet_df in df.items():
                sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)
    return True

def scrub_excel_monitoring(path):
    """Removes duplicates from master_monitoring.xlsx periodic sheets."""
    if not os.path.exists(path):
        return False
    
    wb = openpyxl.load_workbook(path)
    for sheet_name in wb.sheetnames:
        if sheet_name.isdigit() or sheet_name in ["2024", "2025", "2026"]:
            df = pd.read_excel(path, sheet_name=sheet_name)
            keys = ["MERCHANT_GROUP", "DIMENSI", "PERIOD"] # Typical keys
            keys = [k for k in keys if k in df.columns]
            if keys:
                df = df.drop_duplicates(subset=keys, keep='first')
                # Overwrite sheet
                with pd.ExcelWriter(path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
    return True

def scrub_neon_database(engine, schema="public"):
    """
    Performs scrubbing on the Neon PostgreSQL cloud database.
    Uses 'ctid' for targeting duplicates in the absence of a unique ID.
    """
    scrub_targets = {
        "processed_card_monthly": ["merchant_group", "merchant_brand", "transaction_month"],
        "processed_card_history": ["merchant_group", "merchant_brand", "transaction_month"],
        "processed_card_share":   ["merchant_group", "merchant_brand", "transaction_month"],
        "target":                 ["merchant_group", "pm"]
    }

    results = {}
    with engine.begin() as conn:
        for table, keys in scrub_targets.items():
            full_table = f"{schema}.{table}"
            keys_str = ", ".join(keys)
            
            # 1. Deduplicate using ctid
            scrub_sql = text(f"""
                DELETE FROM {full_table}
                WHERE ctid NOT IN (
                    SELECT MIN(ctid)
                    FROM {full_table}
                    GROUP BY {keys_str}
                )
            """)
            try:
                res = conn.execute(scrub_sql)
                results[table] = f"Removed {res.rowcount} duplicates."
            except Exception as e:
                results[table] = f"Error: {str(e)}"

        # 2. Yoshinoya 50x Correction on Neon
        yoshi_sql = text(f"""
            UPDATE {schema}.processed_card_monthly
            SET trx_qris_offus = trx_qris_offus / 50
            WHERE UPPER(merchant_group) = 'YOSHINOYA'
              AND transaction_month LIKE '2025-03%'
              AND trx_qris_offus > 5000000
        """)
        try:
            res_y = conn.execute(yoshi_sql)
            results["yoshinoya_fix"] = f"Normalized {res_y.rowcount} record(s)."
        except Exception as e:
            results["yoshinoya_fix"] = f"Correction failed: {str(e)}"

    return results

def reset_neon_database(engine, schema="public"):
    """
    Purges all data from the Neon PostgreSQL cloud database.
    Truncates all project-relevant tables. Highly Destructive.
    """
    tables_to_clear = [
        "processed_card_monthly", "processed_card_history", "processed_card_share",
        "processed_monitoring",   "processed_monitoring_weekly", "target",
        "mart_merchant_cluster", "raw_master", "raw_card_share", 
        "raw_monitoring", "raw_weekly", "raw_target", "ingestion_runs"
    ]

    results = {}
    with engine.begin() as conn:
        for table in tables_to_clear:
            full_table = f"{schema}.{table}"
            # Check if table exists first to avoid errors
            check_sql = text(f"""
                SELECT EXISTS (
                    SELECT 1 FROM information_schema.tables 
                    WHERE table_schema = :schema AND table_name = :table
                )
            """)
            exists = conn.execute(check_sql, {"schema": schema, "table": table}).scalar()
            
            if exists:
                try:
                    # TRUNCATE is faster and resets identities
                    conn.execute(text(f'TRUNCATE TABLE "{schema}"."{table}" RESTART IDENTITY CASCADE'))
                    results[table] = "Successfully purged."
                except Exception as e:
                    # Fallback to DELETE if truncate fails
                    try:
                        conn.execute(text(f'DELETE FROM "{schema}"."{table}"'))
                        results[table] = "Cleared (via DELETE)."
                    except Exception as e2:
                        results[table] = f"Failed to clear: {str(e2)}"
            else:
                results[table] = "Table not found (Skipped)."

    return results
